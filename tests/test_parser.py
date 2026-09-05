import unittest
from io import BytesIO

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from backend.core.parser import DocumentError, parse_document


class ParserTests(unittest.TestCase):
    def test_markdown_is_parsed_in_memory(self):
        text = parse_document("demo.md", "# Demo\n这是足够长的需求说明和验收标准。".encode())
        self.assertIn("验收标准", text)

    def test_docx_is_parsed_in_memory(self):
        document = Document()
        document.add_paragraph("Demo PRD with enough detail for review and acceptance criteria.")
        buffer = BytesIO()
        document.save(buffer)
        self.assertIn("acceptance", parse_document("demo.docx", buffer.getvalue()))

    def test_docx_tables_are_included(self):
        document = Document()
        document.add_paragraph("Demo PRD with enough introductory detail.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "验收条件"
        table.cell(0, 1).text = "提交后显示成功"
        buffer = BytesIO()
        document.save(buffer)
        parsed = parse_document("demo.docx", buffer.getvalue())
        self.assertIn("验收条件 | 提交后显示成功", parsed)

    def test_text_pdf_is_parsed_in_memory(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 72 720 Td (Demo PRD with acceptance criteria and user value.) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        buffer = BytesIO()
        writer.write(buffer)

        parsed = parse_document("demo.PDF", buffer.getvalue())
        self.assertIn("acceptance criteria", parsed)

    def test_scanned_pdf_without_text_is_rejected_with_ocr_hint(self):
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        buffer = BytesIO()
        writer.write(buffer)

        with self.assertRaisesRegex(DocumentError, "OCR"):
            parse_document("scan.pdf", buffer.getvalue())

    def test_encrypted_pdf_is_rejected(self):
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        writer.encrypt("secret")
        buffer = BytesIO()
        writer.write(buffer)

        with self.assertRaisesRegex(DocumentError, "加密"):
            parse_document("protected.pdf", buffer.getvalue())

    def test_unsupported_extension_lists_pdf(self):
        with self.assertRaisesRegex(DocumentError, r"\.pdf"):
            parse_document("demo.txt", b"A sufficiently detailed product requirement document.")

    def test_oversize_document_is_rejected(self):
        with self.assertRaisesRegex(DocumentError, "3.5MB"):
            parse_document("demo.md", b"a" * 3_500_001)

    def test_excessive_extracted_text_is_rejected(self):
        with self.assertRaisesRegex(DocumentError, "8 万字符"):
            parse_document("demo.md", ("# Demo\n" + "需求" * 40_001).encode())


if __name__ == "__main__":
    unittest.main()
